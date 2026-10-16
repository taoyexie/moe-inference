from pathlib import Path
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import json
import re
from lxml import etree
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

BASE = Path(r'D:\moe inference')
TMP = BASE / 'tmp' / 'patent'
OUT = BASE / 'reports' / 'KITE专利技术交底书_待补实验数据.docx'
REF = Path(r'D:\working\专利技术交底书模板.docx')
SOURCE_HASH = '43ccfd509278626e93f8b5f693dd7248228b6a7c95b7ad907230c7f937e569fe'
assert sha256(REF.read_bytes()).hexdigest() == SOURCE_HASH
TMP.mkdir(parents=True, exist_ok=True)
OUT.parent.mkdir(parents=True, exist_ok=True)

# Black and white patent figures. Diagram labels are identifiers, not claims.
FONT = r'C:\Windows\Fonts\msyh.ttc'
def diagram(path, mode):
    w, h = (1900, 890) if mode == 'system' else (1900, 590)
    im = Image.new('RGB', (w, h), 'white')
    dr = ImageDraw.Draw(im)
    ft = ImageFont.truetype(FONT, 37)
    small = ImageFont.truetype(FONT, 31)
    def box(coords, title, lines=()):
        dr.rectangle(coords, outline='black', width=3)
        x1,y1,x2,y2=coords
        parts=[title]+list(lines)
        gap=56
        top=(y1+y2-gap*(len(parts)-1))/2
        for i,txt in enumerate(parts):
            dr.text(((x1+x2)/2,top+i*gap),txt,fill='black',font=ft if i==0 else small,anchor='mm')
    def arrow(points, label=None):
        dr.line(points,fill='black',width=3)
        x,y=points[-1]; a,b=points[-2]
        if x==a:
            sign=1 if y>b else -1
            dr.polygon([(x,y),(x-10,y-18*sign),(x+10,y-18*sign)],fill='black')
        else:
            sign=1 if x>a else -1
            dr.polygon([(x,y),(x-18*sign,y-10),(x-18*sign,y+10)],fill='black')
        if label:
            dr.text(label[0],label[1],fill='black',font=small,anchor='mm')
    if mode=='system':
        box((35,65,535,285),'101 历史观测模块',('专家与 KV 复用距离直方图','四路径历史计时表'))
        box((35,400,535,560),'102 批次信息模块',('已排程的 Prefill / Decode 构成',))
        box((680,65,1190,285),'103 历史分配模块',('候选容量的 miss 估计','结合路径代价生成历史分配'))
        box((680,400,1190,560),'104 阶段修正模块',('根据当前批次生成修正分配',))
        box((1340,225,1865,470),'105 融合与执行模块',('加权融合并校验容量','输出专家池与 KV 池大小'))
        box((680,680,1865,830),'106 在线更新模块',('回收实际访问、miss 和路径计时记录',))
        arrow([(535,175),(680,175)])
        arrow([(535,480),(680,480)])
        arrow([(1190,175),(1260,175),(1260,285),(1340,285)])
        arrow([(1190,480),(1260,480),(1260,400),(1340,400)])
        arrow([(1600,470),(1600,680)])
        arrow([(680,755),(15,755),(15,175),(35,175)])
    else:
        labels=[('S1',('读取历史与','批次信息')),('S2',('生成历史与','阶段建议')),('S3',('融合并校验','可行容量')),('S4',('调整容量并','执行推理')),('S5',('更新访问与','路径计时'))]
        xs=[30,410,790,1170,1550]
        for x,(num,label) in zip(xs,labels):
            box((x,100,x+320,310),num,label)
        for i in range(4):
            arrow([(xs[i]+320,205),(xs[i+1],205)])
        arrow([(1710,310),(1710,465),(190,465),(190,310)],((940,515),'下一轮使用更新后的历史'))
    im.save(path)

diagram(TMP/'system.png','system')
diagram(TMP/'flow.png','flow')

# Start from the retained template and clone its real paragraph components.
d=Document(str(REF))
source=list(d.paragraphs)
body_proto=deepcopy(source[9]._p)
question_proto=deepcopy(source[8]._p)
section_xml=deepcopy(d.sections[0]._sectPr)
body=d._element.body
for child in list(body):
    if child.tag!=qn('w:sectPr'):
        body.remove(child)

def paragraph(text, role='body', source_index=None):
    from docx.text.paragraph import Paragraph
    el=deepcopy(source[source_index]._p if source_index is not None else question_proto if role in ('question','sub') else body_proto)
    # Remove prior text, bookmark anchors, and stale source identifiers.
    for child in list(el):
        if child.tag!=qn('w:pPr'):
            el.remove(child)
    for attr in list(el.attrib):
        if attr.endswith('paraId') or attr.endswith('textId'):
            del el.attrib[attr]
    body.insert(len(body)-1,el)
    p=Paragraph(el, d._body)
    r=p.add_run('')
    r.font.name='Times New Roman'
    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    r.font.size=Pt(24 if source_index==0 else 14 if source_index==2 else 12)
    r.font.color.rgb=__import__('docx').shared.RGBColor(0,0,0)
    r.bold=role in ('question','sub') or source_index in (0,2,3,4,5,6)
    for part in re.split(r'([BMsNf]_[EKPDhc])',text):
        if not part: continue
        if re.fullmatch(r'[BMsNf]_[EKPDhc]',part):
            for char,subscript in [(part[0],False),(part[2],True)]:
                rr=p.add_run(char);rr._r.insert(0,deepcopy(r._r.rPr));rr.font.subscript=subscript
        else:
            rr=p.add_run(part);rr._r.insert(0,deepcopy(r._r.rPr))
    pf=p.paragraph_format
    if role in ('question','sub'):
        pf.first_line_indent=Cm(0)
        pf.keep_with_next=True
        pf.space_before=Pt(8 if role=='question' else 5)
        pf.space_after=Pt(3)
    elif role=='caption':
        pf.first_line_indent=Cm(0)
        pf.space_after=Pt(5)
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        for rr in p.runs: rr.font.size=Pt(10.5)
    elif source_index is None:
        pf.space_after=Pt(3)
    pf.widow_control=True
    return p

def b(text): return paragraph(text)
def s(text): return paragraph(text,'sub')
def q(i): return paragraph(source[i].text.strip(),'question')
def fig(file,caption):
    p=paragraph('','caption')
    p.paragraph_format.keep_with_next=True
    p.add_run().add_picture(str(TMP/file),width=Cm(16.2))
    docpr=p._p.xpath('.//wp:docPr')[0]
    docpr.set('descr',caption)
    paragraph(caption,'caption')

def table(headers, rows, widths):
    t=d.add_table(rows=1,cols=len(headers))
    t.style='Table Grid'
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    for c,width in zip(t.columns,widths): c.width=Cm(width)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        e=OxmlElement('w:'+edge)
        for k,v in [('val','single'),('sz','6'),('color','D9D9D9')]: e.set(qn('w:'+k),v)
        borders.append(e)
    old=pr.find(qn('w:tblBorders'))
    if old is not None: pr.remove(old)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge,val in [('top','90'),('bottom','90'),('left','110'),('right','110')]:
        e=OxmlElement('w:'+edge);e.set(qn('w:w'),val);e.set(qn('w:type'),'dxa');margins.append(e)
    pr.append(margins)
    for cell,text in zip(t.rows[0].cells,headers): cell.text=text
    for rowdata in rows:
        for cell,text in zip(t.add_row().cells,rowdata): cell.text=text
    for ri,row in enumerate(t.rows):
        rp=row._tr.get_or_add_trPr()
        rp.append(OxmlElement('w:cantSplit'))
        if ri==0: rp.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'F0F0F0');cell._tc.get_or_add_tcPr().append(sh)
            for p in cell.paragraphs:
                pf=p.paragraph_format
                pf.line_spacing=1.2;pf.space_after=Pt(0);pf.space_before=Pt(0);pf.first_line_indent=Pt(0)
                p.alignment=WD_ALIGN_PARAGRAPH.CENTER if ci==0 else WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name='Times New Roman';r.font.size=Pt(10.5);r.bold=ri==0
                    r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')
    paragraph('','caption')

paragraph('专利技术交底书',source_index=0)
paragraph('技 术 交 底 书',source_index=2)
paragraph('发明名称：一种结合历史访问与当前批次阶段信息的混合专家模型联合显存分配方法',source_index=3)
paragraph('技术联系人：待补充',source_index=4)
paragraph('联系人电话：待补充',source_index=5)
paragraph('联系人邮箱：待补充',source_index=6)

q(8)
b('本发明涉及大模型推理系统的内存管理，适用于专家权重与键值缓存不能同时全部常驻 GPU 的混合专家模型推理场景。专家权重用于执行被路由选中的前馈计算，键值缓存用于复用已经处理的上下文；两者共同竞争扣除固定模型部分、激活和临时空间后剩余的显存。给专家池增加容量可能减少专家 miss，却压缩 KV 池并增加 KV 恢复需求；反向调整也具有相同的资源竞争关系。')
b('拟解决的技术问题是：在请求组成动态变化、同类 miss 存在多种处理路径的情况下，如何以较低的在线决策开销，为专家权重和 KV 缓存确定可执行的显存划分，减少由缓存缺失带来的恢复开销。具体包括历史访问统计对当前批次变化响应滞后，以及统一的单次 miss 成本难以区分不同处理路径这两个问题。')
b('本发明采用历史访问和当前批次信息联合决策。历史部分分别维护专家和 KV 的复用距离统计以及四条路径的计时记录，给出考虑恢复代价的分配建议；当前部分读取已经排程但尚未执行的 Prefill 与 Decode 构成，生成阶段修正建议；随后将两种建议融合并转换为满足显存约束的容量。执行结果再反馈到历史记录。阶段信息用于修正资源需求，不通过 GPU 执行时间重新定义工作负载阶段。')

q(10)
s('2.1 相关技术及工作方式')
b('常见的卸载式推理系统将部分专家权重或 KV 缓存保存于 CPU 主存，在 GPU 需要对象时进行搬运或重新计算。固定显存划分简单，但在上下文增长或请求组成变化时，可能使一种资源具有冗余容量而另一种资源频繁发生 miss。仅根据当前 miss 数调整容量也无法直接判断其他候选容量的效果，因为当前 miss 是在当前容量和状态下形成的观测。')
b('WiSP 从工作集角度研究显存受限的 MoE 推理，将专家缓存与 KV 池放在统一可分页预算下，并根据两类资源的容量收益进行联合分配。其重点场景是低并发服务，但论文也涉及启动阶段的 Prefill/Decode 分析与在线控制，不能将其概括为完全未考虑 Prefill。本发明与该类技术的主要区别在于：在历史容量收益分析之外，显式引入下一批次已知的 P/D 构成进行分配修正，并将专家与 KV 的不同恢复路径纳入单次 miss 代价模型。')
b('Fiddler 研究 CPU–GPU 协同执行 MoE 推理，说明专家 miss 不必总是将完整专家权重传入 GPU；也可以在 CPU 上执行专家计算后，仅回传输出。该技术为专家多路径处理提供已有技术基础，但 CPU–GPU 路径选择本身不等于专家池与 KV 池的联合显存分配，也不等于历史分配和当前批次阶段信号的融合。')
s('2.2 与本发明相关的技术不足')
b('第一，仅依赖历史窗口中的访问特征时，若下一批次的请求组成已经变化，分配结果仍可能偏向上一窗口。若改用观测到的 Prefill/Decode GPU 时间划分阶段，则时间又受容量和执行路径影响，阶段判断可能随自身控制动作改变，不能独立反映排程输入。')
b('第二，以每类对象的一个固定 miss 成本作为代理值，不能完整反映路径异构性。固定大小权重上传或 KV 恢复主要受传输影响，而专家 CPU 路径受路由 token 数影响，KV 重计算受恢复范围和上下文影响。在不同运行状态下，较低代价的路径可能不同，单纯减少总 miss 数也未必减少总恢复时间。')
b('第三，将 Prefill 与 Decode 的历史统计完全割裂，容易忽略两者共用物理缓存产生的跨阶段复用。与此同时，复用距离直方图对应的是规定缓存语义下的容量收益；当 CPU 计算等路径绕过缓存准入时，不能无条件把参考 LRU 曲线当成实际缓存 miss 的精确结果。本发明保留统一访问历史，并明确区分预测参考与实际执行状态。')

q(16)
s('3.1 系统组成及总体流程')
b('本方法运行于具有 CPU 主存和 GPU 显存的推理服务器，由推理运行时提供请求阶段、对象访问、路由结果和缓存状态。系统包含历史观测模块 101、批次信息模块 102、历史分配模块 103、阶段修正模块 104、融合与执行模块 105，以及在线更新模块 106。各模块可以作为同一推理服务进程中的软件组件实现，不要求新增专用硬件。')
fig('system.png','图1 历史与当前批次联合驱动的显存分配系统')
b('如图1所示，历史观测模块向历史分配模块提供两类资源的容量收益估计及恢复路径代价；批次信息模块向阶段修正模块提供当前即将执行批次的 P/D 构成。两种分配建议汇入融合与执行模块，该模块输出专家池和 KV 池的容量，而不是重新生成请求或修改用户输入。在线更新模块在实际执行后回收观测结果，供下一轮使用。')
fig('flow.png','图2 联合显存分配及在线更新流程')
b('图2中的 S1 至 S5 分别表示读取历史和批次信息、生成历史与阶段两种建议、融合并投影到可行容量、安全调整并执行推理、更新历史统计。一个决策周期可以是一轮排程，也可以由若干轮组成；同一实现中保持周期和统计窗口定义一致。')

s('3.2 管理对象 容量及记录字段')
b('专家对象由层标识和专家标识确定，包含该专家前馈计算所需的权重矩阵。KV 对象由请求标识、层或层组标识、token 区间和缓存版本确定，包含对应区间的 K/V 张量。对象内容、精度和位置编码信息保持一致，不能将不同请求或不同区间的数据混用。每类资源在一种配置下采用固定对象粒度；两类对象不要求具有相同字节大小。')
b('以 B 表示扣除固定非专家权重、激活、临时工作区和安全预留后的联合可分页显存预算，以 B_E 和 B_K 分别表示专家池与 KV 池容量，满足 B_E 与 B_K 之和不超过 B。以 s_E、s_K 表示两类对象的字节大小，容量对应的对象数分别为 B_E 除以 s_E、B_K 除以 s_K 后向下取整。对齐或取整产生的余量不视为已经缓存了对象。')
b('候选划分同时满足当前不可驱逐对象、传输缓冲区和临时专家执行槽等最低容量约束。对于被正在执行的 kernel 或异步传输引用的对象，设置锁定标记，不能在调整容量时直接释放。若可用显存不足以满足最低约束，维持安全配置或等待资源释放，不强行套用分配结果。')
table(['记录类别','主要字段','作用'],[
    ('访问记录','周期编号、资源类型、对象标识、访问顺序、对象字节数','维护统一资源访问历史及复用距离'),
    ('批次记录','批次编号、Prefill 请求数、Decode 请求数、实际 query token 数','在执行前提供阶段修正信号'),
    ('路径记录','对象类型、处理路径、路由 token 数或上下文长度、完成时间、精度','更新单次 miss 代价模型'),
    ('缓存记录','实际容量、驻留对象、锁定状态、准入及驱逐结果','保证实际状态与执行正确性')
],[2.7,8.2,5.5])

s('3.3 统一历史及候选容量的 miss 预测')
b('分别为专家对象与 KV 对象维护访问顺序和复用距离直方图。Prefill 与 Decode 对同一类资源的访问进入同一条序列，不建立两个互不共享的物理缓存模型。访问指推理执行器对管理对象的逻辑需求；专家侧对应实际非空专家调用，KV 侧对应缓存管理器定义的块读取需求，不能仅依据上下文长度虚构已观测到的物理地址访问。KV 新建写入用于更新驻留状态和参考顺序，不计为需要 Restore/Recompute 的 miss；只有需要复用而未驻留的读取才计入恢复 miss 曲线。')
b('一个对象再次被访问时，统计从其上次访问之后到本次访问之前出现过的不同对象数，记为复用距离 d；首次访问单独记录为冷访问。以 C 表示某类缓存可容纳的对象数，在等大小对象、全相联 LRU 且每次 miss 均准入等条件下，冷访问以及 d 不小于 C 的访问发生 miss。例如序列 A、B、C、A 中，最后一次 A 的复用距离为 2；容量为 2 时该次访问 miss，容量为 3 时该次访问 hit。')
b('统计窗口内，将冷访问数与复用距离不小于候选容量的直方图计数相加，即得该段序列在参考 LRU 语义下的 miss 数。对全部候选容量重复累加即可形成 miss 曲线，不必为每个容量实际创建一个 GPU 缓存。为保留窗口边界前的缓存历史，跨窗口继续维护访问顺序；仅统计计数使用滑动窗口，不能每轮把全部对象重新标为首次访问。')
b('预测未来时，以最近窗口的参考 miss 率或单位访问 miss 数结合历史平均访问量，得到下一决策周期的专家 miss 估计 M_E 和 KV miss 估计 M_K。该外推依赖时间局部性，不读取下一批实际对象标识、实际路由结果或实际 miss。冷启动样本不足时，使用预设安全划分并逐步积累历史。当前批次差异由后续阶段修正处理，而不是在直方图中额外建立 P/D 历史桶。')
b('复用距离曲线是候选容量的参考估计，不等于在所有 fallback 路径下都精确预测物理缓存。采用 CPU 计算后不将专家权重准入 GPU 时，或者恢复对象的附带准入规则不符合 LRU 时，实际状态可能偏离参考状态。因此同时保存实际缓存状态与实际 miss，在线检查预测偏差；偏差持续超过预设阈值时，降低历史建议权重或暂时保留当前安全划分。该处理不改变统一历史与批次修正的主线，但不宣称解决了任意缓存策略下的精确反事实预测。')

s('3.4 四条 fallback 路径及单次 miss 代价模型')
b('miss 次数模型和 cost 模型分别维护。路径模型以单个固定粒度对象为单位，使用历史完成时间标定。阶段名称不是路径 cost 模型的直接输入；专家 CPU 路径主要使用路由到该专家的 token 数 q，KV 重计算主要使用按恢复边界定义的上下文长度 L。对象大小、数值精度、CPU 线程数、传输缓冲方式和重计算前置状态作为配置固定或分别标定。')
table(['对象及路径','处理过程','低开销模型及边界'],[
    ('专家 Upload','主存权重传入 GPU，在 GPU 执行专家计算','固定大小权重的纯上传近似常数；完整恢复还包含 GPU 计算'),
    ('专家 CPU','输入传到 CPU，执行专家计算，输出回传 GPU','在标定范围内按 q 拟合带固定项的线性模型，包含输入输出搬运'),
    ('KV Restore','将主存中的同一 KV 对象传回 GPU','固定大小对象的纯恢复近似常数；要求存在有效主存副本'),
    ('KV Recompute','从有效前置状态重新生成所需 KV 对象','在固定恢复协议及标定范围内按 L 拟合带固定项的线性模型')
],[3.1,6.0,7.3])
b('专家两路径必须从相同任务起点计时，以同一结果在 GPU 可用为终点。完整 Upload 路径包括权重搬运和 GPU 计算，完整 CPU 路径包括输入 D2H、CPU 计算、必要的格式转换和结果 H2D。若分配目标使用额外 miss penalty，则从两条完整路径中统一扣除同任务的常驻专家 GPU 计算基线。此时 Upload 的额外项可近似视为固定上传时间，CPU 路径的额外项在标定范围内拟合为固定项加 q 的比例项，不能将纯上传时间与完整 CPU 路径未经处理直接比较。')
b('KV 两路径恢复同一请求、同一 token 区间、同一层组的 K/V。Restore 的前提是主存副本有效。Recompute 的前提是保存了原 token、位置和恢复所需的前置缓存或检查点，并能按原模型计算出目标对象；缺少前置缓存时须计入实际需要重算的前缀，不能把一层投影时间当作完整恢复时间。恢复协议变化时使用另一组标定参数，不能仅改变 L 而沿用不同协议的系数。')
b('常数和线性模型都是有适用范围的近似。带宽争用、CPU 并行度、矩阵规模、上下文增长和算子切换可能改变曲线，拟合时保留固定开销并记录残差。路径记录表可维护每条路径的样本数、常数或斜率与截距、适用范围和最近更新周期。模型预测超出标定范围或误差过大时，采用保守历史值或分段查表，不预设任意 q 或 L 下都严格线性。')
b('对每类对象，在满足执行前提的路径中估计单次恢复开销较低的路径。历史分配时，尚未产生的 q 采用过去观测的代表值或分布估计；路径实际执行时才允许使用 router 已经产生的真实 q。对于 L，可以使用执行前已经可知的请求上下文信息。多个 q、L 的样本分别求路径开销后再汇总，避免把平均参数上的最小值误当成逐对象选择后的平均开销。')
b('计费不重复计算同一工作。KV 宏重计算中若触发内部专家上传，相关时间包含在该宏路径 cost 内，并以恢复事件标识关联内部操作，不再作为同一决策周期的独立主任务专家恢复项重复相加。一次重算产生多个目标块时，共享的前缀工作只计一次；若逐块独立可加假设不成立，按固定恢复对象组标定，或者将共享恢复作为替代实施方式处理。')

s('3.5 生成基于历史和路径代价的分配建议')
b('从可行专家池容量中生成候选集合，余下预算供 KV 池使用，并按各自对象大小取整。对每个候选划分，从两条 miss 曲线取得 M_E、M_K，结合历史路径记录得到两类对象的平均有效单次 miss 开销。以专家预计 miss 数乘以专家有效开销，加上 KV 预计 miss 数乘以 KV 有效开销，形成该候选划分的估计恢复成本；内部共享恢复按上一节口径扣除重复项。')
b('选取估计恢复成本较低的候选划分作为历史建议。以 f_h 表示历史建议中分给专家池的预算比例，KV 比例为其余部分。若多个候选的差异在预设容差内，可优先保留当前划分以减少迁移。该步骤是在给定历史和简化成本模型上的候选比较，不代表对未知未来请求求得理论全局最优。')

s('3.6 根据当前已排程批次生成阶段修正建议')
b('在本轮计算开始前，从调度器读取已排程批次中 Prefill 请求数 N_P 与 Decode 请求数 N_D。Prefill 指处理新输入或一个已排程的输入分块，Decode 指使用已有上下文生成后续 token 的计算；阶段来自请求状态，而不是通过该轮 GPU 时间占比推断。将 N_P 占两者总数的比例作为当前信号 p，Decode 的比例为剩余部分。空批次不产生新分配。')
b('预先确定 Prefill 倾向的专家比例 f_P 和 Decode 倾向的专家比例 f_D，其中 f_P 大于 f_D，二者均落在可行范围内。按 p 在两个比例之间插值，得到当前批次建议 f_c。直观上，Prefill 请求比例较高时，建议给专家更大容量；Decode 请求比例较高时，建议给 KV 更大容量。这是一种低开销的资源需求修正先验，不断言每个 Prefill 一定更需要专家容量，或者每个 Decode 一定更需要 KV 容量。')
b('f_P、f_D 及信号口径由配置或独立标定确定，评价阶段冻结。采用请求数的实施方式成本低，但一个长 Prefill 与一个短 Prefill 具有相同权重，可能造成误差；替代实施方式可以使用批次内已知 query token 数或其归一化权重。本实施方式不以 GPU 执行时间计算比例，因此不会因分配动作缩短某阶段时间而直接改变该阶段的排程标签。')

s('3.7 融合两种建议并执行安全显存调整')
b('以 λ 表示当前批次建议的权重，取值在 0 到 1 之间。将历史专家比例 f_h 按权重 1−λ、当前专家比例 f_c 按权重 λ 求加权平均，得到融合比例 f。λ 为配置参数，样本充足时按冻结参数运行；冷启动或参考曲线误差过大时可使用事先规定的安全权重。λ 等于 0 表示只使用历史建议，等于 1 表示只使用当前批次建议。')
b('将 f 乘以预算 B 得到专家池的目标字节容量，再将剩余容量分给 KV 池。经过容量边界裁剪、块大小取整和锁定对象检查，形成可执行划分。预算不足、对象对齐或实际锁定状态可能使最终比例不完全等于 f，系统记录目标容量和实际容量，不能只记录加权比例。')
b('在安全同步点应用新容量。缩小某池时，仅驱逐允许移出的对象：专家权重保留可访问的主存来源；KV 若后续还会用到，则保存有效主存副本或保留合法重计算前提。扩容只增加可用槽位，之后由实际访问逐渐填充，不假定扩容立即形成已预热的缓存。若对象仍被 kernel 或传输引用，则延后调整，不释放在用数据。')
b('容量变化很小时可保持现有划分；需要调整时，可以限制每次改变量和最短调整间隔，避免分配器因微小噪声频繁搬运。调整引发的迁移、同步及预热开销单独记录，与 fallback 恢复开销保持一致的评价口径。随后执行已排程批次，发生实际 miss 时根据可用路径、真实 q 和有效恢复状态执行对应路径。')

s('3.8 执行反馈与在线维护')
b('本轮完成后，在线更新模块维护统一资源访问序列、直方图计数和实际 miss 记录，并将每个恢复事件的路径、完成时间、q 或 L 写入路径历史表。传输与计算存在异步提交时，在结果真正可用的完成点结束计时，不能仅用提交函数的返回时间更新 cost。标定和评价中保持精度、缓冲区分配是否计入、CPU 线程数及同步边界一致。')
b('更新只使用已经发生的事件。下一轮先读取冻结的历史版本，再生成建议，避免一个决策周期内边计算边修改其评价依据。主任务访问与 fallback 内部访问使用事件类型区分；实际缓存状态按真实准入和驱逐结果更新，参考 LRU 顺序按参考语义维护，二者不混为一个状态。历史衰减、窗口长度、误差阈值和模型重标定规则由实现配置记录。')

s('3.9 分配机制说明例')
b('以下仅用假设参数说明算法计算步骤，不作为实验数据或性能结论。设可分页预算为 10 GiB，历史模块建议专家占 40%；设置 Prefill 倾向专家占 70%、Decode 倾向专家占 30%。当前批次有 3 个 Prefill 请求和 1 个 Decode 请求，则 Prefill 请求比例为 75%，当前建议为专家占 60%。若当前建议权重 λ 为 0.5，融合后的专家比例为 50%，初步容量为专家 5 GiB、KV 5 GiB，再根据对象大小取整并检查最低容量约束。')
b('下一批次若变成 1 个 Prefill 和 3 个 Decode，在相同倾向配置下，当前建议变为专家占 40%。历史建议不必先等待新的执行时间或 miss 完全体现该变化，分配器已经可以利用排程元数据作出修正。该例说明信息流和计算方式，不证明请求比例足以准确描述全部工作集，也不证明这一调整在所有模型和负载上必然更快。')

s('3.10 实施条件及实验数据预留')
b('一种可采用的实施环境为单 GPU 与 CPU 主存组成的推理服务器。现有可用硬件为 NVIDIA A100 PCIe 40GB、Intel Xeon Gold 5318Y 双路处理器，合计 48 个物理核心和 96 个逻辑线程，主存约 125 GiB。可选择 Qwen3-30B-A3B 等 MoE 模型进行验证；模型、权重精度、对象粒度、软件版本和可分页预算应作为实施例参数固定。本方法并不限定于该型号或模型。')
b('实验数据待补充：真实模型与请求集说明、专家和 KV 访问观测、四路径的单对象计时及拟合误差、相同预算下的分配对照、容量调整开销和必要的正确性检查。真实 trace 驱动的缓存重放与真实端到端推理分别列示；只有完成实际接入后的计时，才填写吞吐、首 token 时延或每 token 时延改善。此处不填写待测数字。')

q(49)
b('针对历史统计响应滞后，本发明将已排程批次构成作为执行前可获得的信号，与历史建议联合使用，使分配器具备提前响应阶段组成变化的机制；不需要等该批次执行完毕后再观察时间占比。')
b('针对统一 miss 成本难以反映路径差异，本发明分别维护四条路径的常数或参数化近似模型，将专家路由 token 数及 KV 恢复上下文等因素用于单次 miss 开销估计，使容量比较不只依赖 miss 总数。模型参数可通过实际完成时间持续维护。')
b('针对阶段历史割裂与缓存语义混淆，本发明保留每类资源跨 P/D 的统一访问历史，同时区分参考 miss 曲线与实际缓存状态；在规定的 LRU 条件下可复用直方图的容量分析能力，在不满足条件时通过误差记录和保守控制限制过度依赖历史估计。')
b('针对显存资源竞争和在线控制开销，本发明在同一字节预算内协调专家池与 KV 池，仅在候选划分之间进行低维比较与融合，并在安全点处理取整、驱逐和扩容。上述优点是由机制产生的预期技术效果，具体改善幅度由后续实测确定。')

q(53)
b('关键技术点一：在专家权重与 KV 缓存共用的显存预算下，把候选容量的 miss 估计与单次 miss 的路径代价估计解耦，再结合形成基于历史的联合分配建议。资源访问历史跨 Prefill/Decode 共享，而 cost 使用与路径计算和搬运相关的参数。')
b('关键技术点二：在推理批次执行前读取排程元数据，以 Prefill/Decode 构成生成当前批次分配建议，并与历史建议加权融合，经可行容量投影后输出唯一的专家池与 KV 池划分。当前建议不依赖已被分配动作改变的 GPU 执行时间。')
b('关键技术点三：将统一访问统计、路径计时、实际缓存状态和安全容量调整连接成在线反馈流程，使下一轮使用已完成事件形成的历史版本，同时明确参考 LRU 预测与真实准入语义之间的边界。')
b('上述技术点应作为协同组合描述。复用距离统计、LRU 规则、CPU–GPU 卸载或某条 fallback 路径单独存在于已有技术中，本交底书不将其单独表述为本发明首次提出。拟保护的重点是其与当前批次阶段修正及联合显存分配的具体组合和信息流。')

q(55)
b('阶段信号的替代：以当前批次中的 query token 数或固定归一化工作量替代请求数，仍从执行前已知元数据生成修正建议。锚点插值可以替换为有限的阶段比例查表，表项由独立标定给出；不需要重新建立按阶段分开的物理缓存。')
b('融合方式的替代：将加权后的目标比例映射到最近可行候选，或只在该比例邻近的候选集合内再次比较成本。误差过大时降低历史权重、暂时保留当前配置或启用安全固定划分。替代方式仍保留历史证据和当前批次信息共同参与决策。')
b('cost 模型的替代：在常数或线性模型不能充分描述实测曲线时，采用分段线性或小规模查表，输入仍为对象大小、q、L 及固定执行配置。对于共享前缀的多个 KV 恢复对象，可采用对象组为粒度的宏恢复计费。')
b('缓存估计的替代：在采用非 LRU、绕过准入或不同大小对象的系统中，可用符合实际准入规则的候选虚拟缓存估计替换参考直方图，或者固定为满足 LRU 前提的执行模式。该替代不改变四路径计时和历史与当前批次融合流程，但具体精度和维护开销需要单独验证。')

q(57)
for txt in [
    'MoE（Mixture of Experts）：混合专家模型。每个 token 通过路由器选择部分专家执行前馈计算，未被选择的专家不参与该 token 的专家运算。',
    'Expert 与 routed token：Expert 指一个专家子网络及其权重；routed token 指在一次专家调用中被路由到该专家的输入 token，数量记为 q。',
    'KV Cache：注意力计算中保存历史 token 的 Key 和 Value 张量的缓存。上下文长度 L 按恢复协议定义为相关历史 token 数，而不是任意输入字符串长度。',
    'Prefill 与 Decode：Prefill 处理新输入 token 或已排程输入分块；Decode 使用已形成的上下文计算后续生成 token。同一 GPU 批次可以包含两类请求。',
    'miss 与 fallback：miss 指所需管理对象在对应 GPU 池中不可用；fallback 指满足该需求的补救路径，本方案包括专家 Upload、专家 CPU 计算后回传、KV Restore 和 KV Recompute。',
    '复用距离与 miss 曲线：复用距离为两次访问同一对象之间访问的不同对象数；直方图统计各距离的出现次数，miss 曲线表示参考缓存容量与 miss 数或 miss 率的关系。',
    'LRU 与准入：LRU 按最近访问顺序驱逐最久未使用对象；准入指 miss 处理后将对象纳入 GPU 缓存。CPU 执行但不上传专家权重属于绕过权重准入。',
    '联合预算与分配比例：B 为扣除固定部分和预留后的可分页字节预算，B_E、B_K 为两池容量；f_h、f_c、f 分别表示历史、当前和融合后的专家预算比例；f_P、f_D 为阶段倾向锚点，λ 为当前建议权重。',
    'Phase-aware 与 path-aware：前者指利用当前排程批次的阶段构成修正资源分配；后者指在代价估计中显式考虑不同恢复路径及其参数。KITE 为本联合分配方法的项目名称。',
]: b(txt)

q(59)
b('[1] Jiamu Zhang, Liang Wu, Mayank Darbari, Liangjie Hong. WiSP: A Working-Set View of Mixture-of-Experts Serving on Extremely Low-Resource Hardware. arXiv:2606.21868, 2026. https://arxiv.org/abs/2606.21868')
b('[2] Keisuke Kamahori, Yile Gu, Kan Zhu, Baris Kasikci. Fiddler: CPU-GPU Orchestration for Fast Inference of Mixture-of-Experts Models. arXiv:2402.07033, 2024. https://arxiv.org/abs/2402.07033')
b('[3] Qwen Team. Qwen3-30B-A3B 模型说明及配置。https://huggingface.co/Qwen/Qwen3-30B-A3B')

# Preserve the template's section geometry and all opaque package parts.
def structure(el):
    return (el.tag,tuple(sorted(el.attrib.items())),el.text,tuple(structure(c) for c in el))
assert structure(d.sections[0]._sectPr)==structure(section_xml)
buf=BytesIO();d.save(buf)
generated=ZipFile(BytesIO(buf.getvalue()))
baseline=ZipFile(REF)
editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml','docProps/core.xml'}
RNS='http://schemas.openxmlformats.org/package/2006/relationships'
rels=etree.fromstring(baseline.read('word/_rels/document.xml.rels'))
old_ids={x.get('Id') for x in rels}
for x in etree.fromstring(generated.read('word/_rels/document.xml.rels')):
    if x.get('Id') not in old_ids:
        rels.append(deepcopy(x))
ct=etree.fromstring(baseline.read('[Content_Types].xml'))
existing={(x.tag,x.get('Extension'),x.get('PartName')) for x in ct}
for x in etree.fromstring(generated.read('[Content_Types].xml')):
    if (x.tag,x.get('Extension'),x.get('PartName')) not in existing:
        ct.append(deepcopy(x))
core=etree.fromstring(baseline.read('docProps/core.xml'))
ns={'dc':'http://purl.org/dc/elements/1.1/','cp':'http://schemas.openxmlformats.org/package/2006/metadata/core-properties'}
for tag,val in [('dc:title','KITE 专利技术交底书'),('dc:creator',''),('cp:lastModifiedBy','')]:
    prefix,local=tag.split(':');node=core.find('{'+ns[prefix]+'}'+local)
    if node is None: node=etree.SubElement(core,'{'+ns[prefix]+'}'+local)
    node.text=val
def xmlbytes(el):return etree.tostring(el,xml_declaration=True,encoding='UTF-8',standalone=True)
parts={'word/document.xml':generated.read('word/document.xml'),
       'word/_rels/document.xml.rels':xmlbytes(rels),
       '[Content_Types].xml':xmlbytes(ct),'docProps/core.xml':xmlbytes(core)}
with ZipFile(OUT,'w',ZIP_DEFLATED) as z:
    for info in baseline.infolist(): z.writestr(deepcopy(info),parts.get(info.filename,baseline.read(info.filename)))
    for name in generated.namelist():
        if name.startswith('word/media/') and name not in baseline.namelist():
            z.writestr(name,generated.read(name))

final=ZipFile(OUT)
preserved=[]
for name in baseline.namelist():
    if name not in editable:
        assert baseline.read(name)==final.read(name),name
        preserved.append({'part':name,'bytes':len(baseline.read(name)),'sha256':sha256(baseline.read(name)).hexdigest()})
assert sha256(REF.read_bytes()).hexdigest()==SOURCE_HASH
check=Document(str(OUT))
texts='\n'.join(p.text for p in check.paragraphs)
for word in ['LED','机器人','16.7%','270.5%','102.5%']:
    assert word not in texts,word
for qi in [8,10,16,49,53,55,57,59]:
    assert source[qi].text.strip() in texts
report={'output':str(OUT),'paragraphs':len(check.paragraphs),'tables':len(check.tables),
        'figures':len(check.inline_shapes),'text_chars':len(texts),
        'preserved_parts':preserved,'source_sha256':SOURCE_HASH,
        'rendering':'Word export and bundled Poppler after packaged renderer unavailable'}
(TMP/'structural_qa.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='preserved_parts'},ensure_ascii=False,indent=2))
