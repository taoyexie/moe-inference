"""Describe the pending WiSP-style comparison; keep old measurements accurately labeled."""
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from update_patent_baseline import media, rewrite


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'reports/KITE专利技术交底书_终稿_baseline说明版.docx'
OUTPUT = ROOT/'reports/KITE专利技术交底书_WiSP对照待测版.docx'
QA = ROOT/'tmp/patent_wisp_20261016'

COMPARISON = [
    '本次负载共 16 个计算批次：首批包含四条 Prefill，后续 15 批包含四条 Decode。'
    '后续实验采用 WiSP 式固定成本分配作为 baseline，与完整 KITE 比较。'
    '该对照依据 WiSP 第 3.3 节式（2）的两类固定 miss 代价目标构建，'
    '用于比较分配策略，不等同于原始 WiSP 在其服务引擎上的完整复现。',
    'WiSP 式 baseline 的成本模型：Expert miss 使用一个固定专家上传代价，'
    'KV miss 使用一个固定块重计算代价；分配评分为 Expert 预测 miss 数乘以 Expert 固定代价，'
    '加上 KV 预测 miss 数乘以 KV 固定代价。两类常数拟在工作负载启动阶段标定并冻结。'
    '该评分不再比较四条 fallback 路径，也不随 routed-token 数 q 或上下文长度调整单次代价。',
    'WiSP 式 baseline 的分配方式：依据历史复用距离统计估计候选容量下的 miss，'
    '在相同可分页预算及预先规定的容量约束下，选择固定成本评分最低的可行 Expert/KV 划分。'
    '本对照配置拟与 KITE 一样每 16 批刷新历史建议，但不加入当前批次 P/D 请求比例修正。'
    '固定的是两类 miss 的成本系数，不是显存比例；历史改变时分配仍可改变。',
    '完整 KITE 将 miss 统计与路径成本分开维护：Upload 和 Restore 使用常数，'
    'Expert CPU 和 KV 前缀重放使用带截距的线性模型，并以 0.65 权重融合当前 P/D 信号，'
    '选取最近可行划分。后续两组对照固定相同的模型、请求、调度、精度、可分页预算、'
    '输出长度和计时边界，分别记录分配策略与实际恢复路径，比较完整工作负载完成时间。',
]

PENDING_NOTE = (
    'WiSP 对照尚未实测。表2及图3至图5暂保留原四路径固定成本 baseline 与 KITE 的实测结果，'
    '其中 Fixed-cost baseline 图例仍指原对照，不代表 WiSP。以下数值和改善幅度仅适用于原对照；'
    '完成 WiSP 式对照测试后，再统一替换结果表、图片及对应技术效果数值。'
)


def one(doc, prefix):
    matches = [p for p in doc.paragraphs if p.text.startswith(prefix)]
    assert len(matches) == 1, (prefix, len(matches))
    return matches[0]


def main():
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    doc = Document(SOURCE)
    prefixes = ['本次负载共 16 个计算批次：', 'baseline 的成本模型：',
                'baseline 的分配方式：', '完整 KITE 在相同执行框架上，']
    for prefix, text in zip(prefixes, COMPARISON):
        rewrite(one(doc, prefix), text)

    timing = one(doc, '在每个预算下，两种配置各执行一次完整工作负载。')
    rewrite(timing, '以下保留结果的原实验中，' + timing.text.replace('baseline', '原 baseline'))
    rewrite(one(doc, '实测结果'), '已有对照结果与 WiSP 待测说明')
    caption = one(doc, '表2 两种预算下的完整工作负载实测结果')
    note = caption.insert_paragraph_before()
    body_format = one(doc, '本次负载共 16 个计算批次：')._p.pPr
    if body_format is not None:
        note._p.insert(0, deepcopy(body_format))
    note.add_run()
    note.runs[0]._r.insert(0, deepcopy(timing.runs[0]._r.rPr))
    rewrite(note, PENDING_NOTE)
    rewrite(caption, '表2 原四路径固定成本 baseline 与 KITE 的实测结果')
    # Keep all numbers intact and name the old comparison explicitly at point of use.
    for prefix in ('如图3所示，', '如图5所示，'):
        paragraph = one(doc, prefix)
        text = paragraph.text.replace('baseline', '原 baseline')
        if prefix == '如图3所示，':
            text = text.replace('6 GiB 预算下，', '与原 baseline 比较，6 GiB 预算下，', 1)
        rewrite(paragraph, text)
    effect = one(doc, '针对显存资源竞争和在线控制开销，')
    rewrite(effect, effect.text.replace('第3.10节实施例中，', '第3.10节已测的原对照实施例中，')
            .replace('这些结果体现了', 'WiSP 对照结果待补。上述原对照结果体现了'))
    for cell in doc.tables[3].rows[0].cells:
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.text = run.text.replace('baseline', '原 baseline')
    doc.save(OUTPUT)

    old, new = Document(SOURCE), Document(OUTPUT)
    assert len(old.inline_shapes) == len(new.inline_shapes) == 5
    assert len(old.tables) == len(new.tables) == 4
    assert media(SOURCE) == media(OUTPUT), 'Every embedded image must remain unchanged.'
    old_numbers = [[c.text for c in r.cells] for r in old.tables[3].rows[1:]]
    new_numbers = [[c.text for c in r.cells] for r in new.tables[3].rows[1:]]
    assert old_numbers == new_numbers
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == source_hash
    with ZipFile(OUTPUT) as archive:
        text = archive.read('word/document.xml').decode('utf-8')
        assert '基线' not in text and 'WiSP 对照尚未实测' in text
    QA.mkdir(parents=True, exist_ok=True)
    manifest = dict(source=str(SOURCE), source_sha256=source_hash, output=str(OUTPUT),
                    output_sha256=hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
                    comparison='pending WiSP-style two fixed costs, history-only allocation',
                    old_results_retained_and_labeled=True, images_byte_identical=True,
                    results_table_numbers_unchanged=True, visual_review='pending')
    (QA/'build_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
