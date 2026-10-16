"""Complete the retained KITE disclosure with audited measurements and original bars.

The supplied draft and all raw experiment records remain unchanged. No new GPU runs.
"""
import argparse
from copy import deepcopy
import hashlib
import json
from pathlib import Path
from statistics import mean
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm,Pt,RGBColor

from audit_kite_parallel_decode import audit

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT/'reports/KITE专利技术交底书_待补实验数据.docx'
RUN = ROOT/'results/parallel-decode-search-20261016T132937Z-8668e263'
FIGURES = ROOT/'reports/kite_parallel_decode_20261016/figures'
OUTPUT = ROOT/'reports/KITE专利技术交底书_终稿.docx'
QA = ROOT/'tmp/patent_final_20261016'


def run_font(run,size=12,bold=False):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor(0,0,0)
    fonts = run._r.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:ascii'),'Times New Roman')
    fonts.set(qn('w:hAnsi'),'Times New Roman')
    fonts.set(qn('w:eastAsia'),'宋体')


def replace_text(paragraph,text,bold=False,size=12):
    for child in list(paragraph._p):
        if child.tag!=qn('w:pPr'):
            paragraph._p.remove(child)
    run_font(paragraph.add_run(text),size,bold)


def before(anchor,text='',kind='body'):
    p = anchor.insert_paragraph_before()
    p.style = 'Normal'
    p.paragraph_format.line_spacing = 1.5 if kind=='body' else 1.15
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5 if kind=='body' else 6)
    p.paragraph_format.first_line_indent = Pt(24 if kind=='body' else 0)
    p.paragraph_format.widow_control = True
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY if kind=='body' else WD_ALIGN_PARAGRAPH.LEFT
    if kind=='subhead':
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(8)
    if kind=='caption':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if text:
        run_font(p.add_run(text),10.5 if kind=='caption' else 12,kind=='subhead')
    return p


def table_before(doc,anchor,rows,widths):
    table = doc.add_table(rows=1,cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col,width in zip(table.columns,widths):
        col.width = Cm(width)
    for row_index,values in enumerate(rows):
        row = table.rows[0] if row_index==0 else table.add_row()
        props = row._tr.get_or_add_trPr()
        no_split = OxmlElement('w:cantSplit')
        props.append(no_split)
        if row_index==0:
            repeat = OxmlElement('w:tblHeader')
            props.append(repeat)
        for i,(cell,text,width) in enumerate(zip(row.cells,values,widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for name,value in [('top',80),('bottom',80),('left',100),('right',100)]:
                item = OxmlElement('w:'+name)
                item.set(qn('w:w'),str(value))
                item.set(qn('w:type'),'dxa')
                margins.append(item)
            tc.append(margins)
            if row_index==0:
                shading = OxmlElement('w:shd')
                shading.set(qn('w:fill'),'E7EBEF')
                tc.append(shading)
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index==0 or (len(widths)>2 and i>0) else WD_ALIGN_PARAGRAPH.LEFT
            run_font(p.add_run(text),10.5,row_index==0)
    borders = OxmlElement('w:tblBorders')
    for name in ('top','left','bottom','right','insideH','insideV'):
        edge = OxmlElement('w:'+name)
        edge.set(qn('w:val'),'single')
        edge.set(qn('w:sz'),'6')
        edge.set(qn('w:color'),'D9D9D9')
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    anchor._p.addprevious(table._tbl)
    before(anchor,'',kind='caption')
    return table


def figure_before(anchor,filename,caption):
    image = before(anchor,kind='image')
    image.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image.paragraph_format.keep_with_next = True
    image.paragraph_format.keep_together = True
    image.paragraph_format.space_before = Pt(5)
    image.paragraph_format.space_after = Pt(0)
    shape = image.add_run().add_picture(str(FIGURES/filename),width=Cm(14.3))
    shape._inline.docPr.set('descr',caption)
    before(anchor,caption,kind='caption')


def complete(replace_own=False):
    if OUTPUT.exists() and not replace_own:
        raise FileExistsError('Refusing to overwrite an existing final document.')
    verified = audit(RUN)
    source_hash = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    data = json.loads((FIGURES/'figure_data.json').read_text(encoding='utf-8'))
    assert data['audited'] and len(data['rows'])==4
    measured = {(int(r['budget_gib']),r['mode']):r for r in data['rows']}
    doc = Document(SOURCE)
    original = list(doc.paragraphs)
    if 'Title' not in doc.styles:
        title_style = doc.styles.add_style('Title',WD_STYLE_TYPE.PARAGRAPH)
        title_style.base_style = doc.styles['Normal']
    original[0].style = doc.styles['Title']
    original[0].paragraph_format.keep_with_next = True
    original[0].paragraph_format.space_after = Pt(10)
    original[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    replace_text(original[0],'专利技术交底书',bold=True,size=22)
    # Remove the duplicated title line, not the invention name or technical text.
    if original[1].text=='技 术 交 底 书':
        original[1]._p.getparent().remove(original[1]._p)
    for index in (3,4,5):
        replace_text(original[index],original[index].text.replace('待补充','________________'),bold=True)
    replace_text(original[65],'3.10 具体实施环境与实验结果',bold=True)
    replace_text(original[66],
        '本实施例在单 GPU 与 CPU 主存组成的推理服务器上实现完整 KITE 执行链，使用 Qwen3-30B-A3B 的完整 BF16 权重、真实路由器和自回归生成。专家权重的主副本保存在 CPU 主存，固定模型部分位于 GPU；专家池和 KV 池在同一可分页预算内分配。硬件型号用于说明本实施例，本方法不限定于该型号或模型。')
    original[67]._p.getparent().remove(original[67]._p)
    anchor = original[68]

    before(anchor,'实施配置',kind='subhead')
    p = before(anchor,'表1 服务器及运行配置',kind='caption')
    p.paragraph_format.keep_with_next = True
    table_before(doc,anchor,[
        ['项目','实施配置'],
        ['GPU 与主存','NVIDIA A100 PCIe 40GB 单卡；主存约 125 GiB'],
        ['CPU','双路 Intel Xeon Gold 5318Y；合计 48 个物理核心、96 个逻辑线程；本实验固定使用 8 个 CPU 线程'],
        ['软件及精度','Python 3.12.3；PyTorch 2.7.1+cu118；Transformers 4.51.3；BF16 模型权重与 GPU 计算；CPU 专家计算使用 FP32 副本'],
        ['管理对象','单个专家权重约 9 MiB；KV 页包含 128 tokens 的全 48 层 K/V，单页 12 MiB；共享显存以 3 MiB 帧管理'],
        ['可分页预算','分别设为 6 GiB 和 7 GiB；固定权重、激活及临时工作区另行记账，显存保留安全余量'],
        ['历史及当前信号','每 16 批刷新历史曲线；当前建议权重为 0.65；Prefill 与 Decode 的专家份额锚点分别为 0.8 和 0.2'],
        ['候选划分及副本缓存','专家份额候选为 0.1、0.25、0.5、0.75、0.9，投影到最近可行容量；CPU FP32 专家副本缓存上限为 8 GiB']
    ],[3.1,13.1])
    before(anchor,'可分页预算是人为设置的缓存压力条件，不表示 A100 的全部 40GB 显存均可用于两类对象，也不表示设备总显存自然耗尽。对象粒度、精度、线程数及执行器在两组之间一致。')
    before(anchor,'每次完整工作负载在开始阶段执行四条路径的成本标定。专家路径使用 q 为 1、8、32 的样本，KV 路径使用上下文长度为 128、512 的样本；以实际完成时间建立常数或带截距的线性近似。没有 FP32 副本时计入冷转换开销。超出标定范围的情况保守采用专家上传或 KV 恢复，不静默外推。KV 备用重计算从原始及已生成 token 重放前缀，多块共享同一重放任务时只执行和计费一次。')

    before(anchor,'请求负载与对照方法',kind='subhead')
    before(anchor,'构造四条同时到达的请求，输入长度依次为 8192、11264、11264、11264 tokens，每条生成 16 个真实输出 tokens，总输出为 64 tokens。提示正文为留档的手写文本，经重复和截断达到指定长度。使用真实路由与贪心自回归生成，统一忽略 EOS 提前终止以固定输出量，不人为指定专家或 routed-token 数。四条请求的位置和 KV 独立，注意力相互隔离；不使用 chunked prefill、答案缓存或前缀缓存。')
    before(anchor,'本次负载共 16 个计算批次：首批包含四条 Prefill，后续 15 批包含四条 Decode。比较固定成本基线和完整 KITE 两种配置。固定成本基线使用同一启动样本的常数统计值，按历史曲线选择划分；KITE 使用常数或线性路径模型，并将历史建议与当前 P/D 请求比例修正融合。两组共用执行器、请求、预算、精度、fallback 集合和初始空 GPU 缓存。该基线是本实施例的固定成本历史分配对照，不作为原始 WiSP 实现。')
    before(anchor,'在每个预算下，两种配置各执行一次完整工作负载。6 GiB 下先执行基线、后执行 KITE；7 GiB 下交换两组顺序。完成时间包含各组自己的启动标定、池初始化、全部 Prefill/Decode、在线统计和决策、容量调整、实际搬运、写回及过期释放。模型加载在工作负载开始前完成，时间单独记录，不计入图3的工作负载完成时间。TPOT 按每条请求首个至末个输出的时间差除以 15 后，在同一负载的四条请求之间取均值。')

    before(anchor,'实测结果',kind='subhead')
    p = before(anchor,'表2 两种预算下的完整工作负载实测结果',kind='caption')
    p.paragraph_format.keep_with_next = True
    def row(metric,label,digits):
        return [label]+[f'{measured[(b,m)][metric]:.{digits}f}' if digits else str(measured[(b,m)][metric])
            for b,m in [(6,'baseline'),(6,'kite'),(7,'baseline'),(7,'kite')]]
    main = {(r['budget_mib']//1024,r['mode']):r for r in json.loads((RUN/'summary.json').read_text(encoding='utf-8'))['tasks']}
    table_before(doc,anchor,[
        ['指标','6 GiB\n基线','6 GiB\nKITE','7 GiB\n基线','7 GiB\nKITE'],
        row('e2e_s','完整完成时间 s',3),
        ['其中主执行时间 s']+[f'{main[(b,m)]["execution_inclusive_ms"]/1000:.3f}' for b,m in [(6,'baseline'),(6,'kite'),(7,'baseline'),(7,'kite')]],
        row('tpot_s','TPOT s/token',3),
        row('kv_restores','实际 KV 恢复次数',0)
    ],[4.5,2.95,2.95,2.95,2.95])
    before(anchor,'如图3所示，6 GiB 预算下，完整工作负载耗时由 82.174 秒降至 75.348 秒，降低 8.31%，对应加速比约 1.091；7 GiB 下由 78.091 秒降至 75.634 秒，降低 3.15%。两组实际输出数量一致。6 GiB 的主执行时间由 51.646 秒降至 48.006 秒，表明收益不只发生在启动区间。')
    figure_before(anchor,'01_e2e_latency.png','图3 不同可分页预算下的完整工作负载完成时间')
    before(anchor,'如图4所示，6 GiB 下 TPOT 由约 1.535 秒降至 1.403 秒，降低 8.59%；7 GiB 下由约 1.541 秒降至 1.392 秒，降低 9.68%。图中的数值为同一工作负载内四条请求的 TPOT 均值，不是多次独立试验的均值。')
    figure_before(anchor,'02_tpot.png','图4 不同可分页预算下的输出 token 间隔')
    before(anchor,'如图5所示，6 GiB 下实际 KV 恢复次数由 4934 次降至 200 次，减少 95.95%；7 GiB 下由 4934 次降至 179 次，减少 96.37%。所示次数来自实际执行的恢复事件，而不是复用距离曲线推算值。6 GiB 基线的 KV 容量始终为 256 页，KITE 在 Decode 阶段扩为 384 页；本次 KV 峰值为 332 页，因此 KITE 可以覆盖该工作集而基线不能。7 GiB 下对应容量为基线 298 页和 KITE Decode 448 页。')
    figure_before(anchor,'03_kv_restores.png','图5 不同可分页预算下的实际 KV 恢复次数')
    before(anchor,'本次两种配置的 Expert GPU 命中数均为零，缩小专家池没有损失已经存在的专家命中，从而避免了以增加专家恢复抵消 KV 收益的情况。日志记录的 6 GiB 实际 KV 恢复累计时间由约 2.884 秒降至 0.115 秒。以上实测支持完整策略在该长输入、短输出、四请求受控场景中的整体效果，不将全部完成时间差单独归因于某一模块，也不据此推断任意请求或长生成稳态下具有相同增益。')
    before(anchor,'显存、主存及恢复事件均记录于运行日志。6 GiB 下两组 GPU 分配峰值均约 21.49 GiB，主机 KV 峰值为 3984 MiB，CPU FP32 副本峰值为 126 MiB，满足所设预算和安全余量；运行期间未观察到其他 GPU 计算进程干扰。')

    replace_text(original[72],
        '针对显存资源竞争和在线控制开销，本发明在同一字节预算内协调专家池与 KV 池，仅在候选划分之间进行低维比较与融合，并在安全点处理取整、驱逐和扩容。第3.10节实施例中，6 GiB 可分页预算下完整工作负载耗时降低 8.31%，TPOT 降低 8.59%，实际 KV 恢复次数减少 95.95%。这些结果体现了整体方案在所述受控场景下的技术效果，具体改善幅度随工作负载及系统配置变化。')

    # Native Title is black and has no theme rule. Existing body/table/diagram styles stay intact.
    title = doc.styles['Title']
    title.font.color.rgb = RGBColor(0,0,0)
    for paragraph in doc.paragraphs:
        if paragraph.style.name=='Title':
            borders = paragraph._p.get_or_add_pPr().find(qn('w:pBdr'))
            if borders is not None:
                borders.getparent().remove(borders)
    doc.core_properties.title = 'KITE 专利技术交底书'
    doc.core_properties.subject = '历史访问与当前批次阶段信息联合驱动的混合专家模型显存分配'
    doc.core_properties.comments = ''
    OUTPUT.parent.mkdir(parents=True,exist_ok=True)
    doc.save(OUTPUT)
    final = Document(OUTPUT)
    text = '\n'.join(p.text for p in final.paragraphs)
    assert '实验数据待补充' not in text and '实验数据预留' not in text and '待补充' not in text
    assert len(final.inline_shapes)==5
    assert len(final.tables)==4
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==source_hash
    image_hashes = {}
    with ZipFile(OUTPUT) as archive:
        members = {hashlib.sha256(archive.read(n)).hexdigest() for n in archive.namelist() if n.startswith('word/media/')}
        for name in ('01_e2e_latency.png','02_tpot.png','03_kv_restores.png'):
            digest = hashlib.sha256((FIGURES/name).read_bytes()).hexdigest()
            assert digest in members,'Original first-version chart was not embedded.'
            image_hashes[name] = digest
    QA.mkdir(parents=True,exist_ok=True)
    evidence = dict(source=str(SOURCE),source_sha256=source_hash,output=str(OUTPUT),
        output_sha256=hashlib.sha256(OUTPUT.read_bytes()).hexdigest(),
        original_figures_preserved=True,experimental_chart_version='original zero-baseline bars',
        embedded_experimental_image_sha256=image_hashes,images=len(final.inline_shapes),tables=len(final.tables),
        original_structure_preserved=True,raw_experiment_audit=verified['status'])
    (QA/'build_manifest.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps(evidence,ensure_ascii=False,indent=2))


if __name__=='__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--replace-own-output',action='store_true')
    complete(parser.parse_args().replace_own_output)
